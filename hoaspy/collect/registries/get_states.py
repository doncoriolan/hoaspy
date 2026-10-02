#!/usr/bin/env python
"""Collect state-level HOA registries and corporate registries for every
state configured in `state_sources.yml`, and keep `coverage.json` current.

    ./venv/bin/python -m hoaspy.collect.registries.get_states                # everything configured
    ./venv/bin/python -m hoaspy.collect.registries.get_states --state CO NV  # specific states

Source kinds (per state, in state_sources.yml):
    socrata_corp      corporate registry on a Socrata portal -> state_corps.jsonl
    socrata_registry  HOA registry on a Socrata portal       -> state_registries.jsonl
    csv_corp          direct CSV/TSV download                -> state_corps.jsonl
    csv_registry      direct CSV/TSV download                -> state_registries.jsonl
                      (`constants:` fills columns the file lacks, e.g. county)
    arcgis_corp       corporate registry on an ArcGIS FeatureServer -> state_corps.jsonl
    tecuity_corp      Tecuity business-search API (ID SOSBiz, ND FirstStop) -> state_corps.jsonl
                      (search sweep, so coverage is recorded as partial)

Corporate rows are filtered to association-looking names (server-side for
Socrata, locally for CSVs). Registry rows are kept whole — the registry
itself is the filter.

Output: records/state_corps.jsonl, records/state_registries.jsonl — one row
per record, unified schema, merged across states (a re-run replaces that
state's rows only).
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import logging
import re
import sys
import time
import zipfile
from datetime import datetime
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import requests
import yaml

from hoaspy.collect.registries import records_socrata
from hoaspy.collect.registries import records_tecuity
from hoaspy import ROOT
SOURCES = ROOT / "state_sources.yml"
COVERAGE = ROOT / "coverage.json"
OUT_DIR = ROOT / "records"
CORPS = OUT_DIR / "state_corps.jsonl"
REGS = OUT_DIR / "state_registries.jsonl"

USER_AGENT = "macos:hoa-public-records:0.1 (academic research)"

ASSOC_RE = re.compile(
    r"(HOMEOWNER|HOME OWNER|CONDOMINIUM|\bCONDO\b|PROPERTY OWNERS|"
    r"COMMUNITY ASSOCIATION|MASTER ASSOCIATION|OWNERS ASSOCIATION|"
    r"RESIDENTS ASSOCIATION|TOWNHOME|TOWNHOUSE|\bHOA\b)", re.I)

log = logging.getLogger("states")


def _get(row: dict, spec) -> str:
    if not spec:
        return ""
    if isinstance(spec, list):
        return " ".join(str(row.get(c) or "").strip() for c in spec).strip()
    return str(row.get(spec) or "").strip()


def _units(v: str) -> int | None:
    """A unit count as a number; None when the column is empty or not one."""
    digits = v.replace(",", "").strip()
    return int(digits) if digits.isdigit() else None


def _iso_date(v: str) -> str:
    """Normalize '3/3/2004' / '7/20/1987 12:00:00 AM' / ISO-ish to YYYY-MM-DD."""
    v = (v or "").split()[0] if (v or "").split() else ""
    for fmt in ("%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(v, fmt).date().isoformat()
        except ValueError:
            pass
    return v[:10]


def _iter_rows(content: bytes, cfg: dict):
    if cfg.get("xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it)]
        for r in it:
            yield {h: ("" if v is None else str(v)) for h, v in zip(header, r)}
        return
    text = content.decode(cfg.get("encoding", "utf-8"), "replace")
    yield from csv.DictReader(io.StringIO(text), delimiter=cfg.get("delimiter", ","))


def fetch_csv(session: requests.Session, cfg: dict, state: str,
              registry: bool) -> list[dict]:
    contents: list[bytes] = []
    if cfg.get("local_glob"):
        paths = sorted(ROOT.glob(cfg["local_glob"]))
        if not paths:
            raise FileNotFoundError(f"no files match {cfg['local_glob']}")
        contents = [p.read_bytes() for p in paths]
    else:
        local = cfg.get("local_cache")
        if local and Path(local).exists():
            contents = [Path(local).read_bytes()]
        else:
            resp = session.get(cfg["url"], timeout=600)
            resp.raise_for_status()
            content = resp.content
            if local:
                Path(local).write_bytes(content)
            contents = [content]
    if cfg.get("zip"):
        z = zipfile.ZipFile(io.BytesIO(contents[0]))
        member = next(n for n in z.namelist() if n.lower().endswith((".csv", ".txt")))
        contents = [z.read(member)]
    fields = cfg["fields"]
    exclude_status = re.compile(cfg["exclude_status"], re.I) if cfg.get("exclude_status") else None
    officials = load_officials(cfg.get("officials"))
    out = []
    seen: set[str] = set()
    for content in contents:
        for row in _iter_rows(content, cfg):
            for k, v in (cfg.get("constants") or {}).items():
                row.setdefault(k, v)   # e.g. county for a single-county list
            name = " ".join(_get(row, fields.get("name")).split())
            if not name:
                continue
            if not registry and not ASSOC_RE.search(name):
                continue
            status = _get(row, fields.get("status"))
            if exclude_status and exclude_status.search(status):
                continue
            rid = _get(row, fields.get("record_id"))
            if cfg.get("dedupe") and rid:
                if rid in seen:
                    continue
                seen.add(rid)
            out.append({
                "state": state,
                "source": cfg.get("source", cfg.get("url", "")),
                "source_url": cfg.get("page", cfg.get("url", "")),
                "record_id": rid,
                "name": name,
                "corp_status": status,
                "status": status,
                "address": _get(row, fields.get("address")),
                "city": _get(row, fields.get("city")).title(),
                "county": _get(row, fields.get("county")).title(),
                "zip": _get(row, fields.get("zip"))[:5],
                "incorporated": _iso_date(_get(row, fields.get("incorporated"))),
                "registered_agent": " ".join(_get(row, fields.get("agent")).split()),
                "entity_type": _get(row, fields.get("type")),
                "units": _units(_get(row, fields.get("units"))),
                "manager_name": _get(row, fields.get("manager")),
                "officers": officials.get(rid, []),
            })
    return out


def load_officials(cfg: dict | None) -> dict[str, list[dict]]:
    """Optional companion CSV of officers keyed by the parent record id
    (e.g. Alaska DCCED's OfficialsDownload.csv). Returns
    {record_id: [{"name": ..., "title": ...}, ...]} in the same shape
    Sunbiz officers use, so build_site/app.js render them unchanged.
    A missing file is not an error — the enrichment is simply skipped."""
    if not cfg:
        return {}
    path = Path(cfg["local_cache"])
    if not path.exists():
        log.warning("officials file %s missing — skipping officer enrichment", path)
        return {}
    out: dict[str, list[dict]] = {}
    seen: set[tuple[str, str, str]] = set()
    with path.open(newline="", encoding="utf-8-sig") as fh:
        for row in csv.DictReader(fh):
            rid = (row.get(cfg["key"]) or "").strip()
            last = " ".join((row.get(cfg["last"]) or "").split())
            first = " ".join((row.get(cfg.get("first", "")) or "").split())
            title = " ".join((row.get(cfg.get("title", "")) or "").split())
            name = f"{first} {last}".strip()
            if not rid or not name or (rid, name.upper(), title) in seen:
                continue
            seen.add((rid, name.upper(), title))
            out.setdefault(rid, []).append({"name": name, "title": title})
    log.info("officials: %d officer rows for %d records from %s",
             sum(len(v) for v in out.values()), len(out), path)
    return out


def _epoch_ms(v) -> str:
    try:
        return datetime.utcfromtimestamp(int(v) / 1000).date().isoformat()
    except (TypeError, ValueError, OSError):
        return str(v or "")[:10]


def fetch_arcgis(session: requests.Session, cfg: dict, state: str) -> list[dict]:
    """ArcGIS FeatureServer registry (currently DC). Server-side name filter,
    resultOffset paging."""
    fields = cfg["fields"]
    patterns = ["HOMEOWNER", "CONDOMINIUM", "PROPERTY OWNERS", "OWNERS ASSOCIATION",
                "COMMUNITY ASSOCIATION", "MASTER ASSOCIATION", "TOWNHOME", "TOWNHOUSE"]
    where = " OR ".join(f"UPPER({fields['name']}) LIKE '%{p}%'" for p in patterns)
    out: list[dict] = []
    offset = 0
    while True:
        resp = session.get(cfg["url"], params={
            "where": where, "outFields": "*", "f": "json",
            "resultOffset": offset, "resultRecordCount": 2000,
            "orderByFields": fields.get("record_id", "OBJECTID")}, timeout=120)
        resp.raise_for_status()
        d = resp.json()
        if "error" in d:
            raise RuntimeError(f"arcgis error: {d['error']}")
        feats = d.get("features", [])
        for f in feats:
            a = f["attributes"]
            out.append({
                "state": state,
                "source": cfg.get("source", cfg["url"]),
                "source_url": cfg.get("page", cfg["url"]),
                "record_id": str(a.get(fields.get("record_id")) or ""),
                "name": " ".join(str(a.get(fields.get("name")) or "").split()),
                "corp_status": str(a.get(fields.get("status")) or ""),
                "address": str(a.get(fields.get("address")) or ""),
                "city": str(a.get(fields.get("city")) or "").title(),
                "county": "",
                "zip": str(a.get(fields.get("zip")) or "")[:5],
                "incorporated": _epoch_ms(a.get(fields.get("incorporated"))),
                "registered_agent": " ".join(str(a.get(fields.get("agent")) or "").split()),
                "entity_type": str(a.get(fields.get("type")) or ""),
            })
        offset += len(feats)
        log.info("%s arcgis: %d rows so far", state, offset)
        if len(feats) < 2000 and not d.get("exceededTransferLimit"):
            break
        time.sleep(0.5)
    return out


def load_existing(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open() as fh:
        return [json.loads(line) for line in fh]


def write_merged(path: Path, fresh: list[dict], states_replaced: set[str]) -> int:
    keep = [r for r in load_existing(path) if r["state"] not in states_replaced]
    merged = keep + fresh
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    with tmp.open("w") as fh:
        for r in merged:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    tmp.replace(path)
    return len(merged)


def update_coverage(state: str, key: str, info: dict) -> None:
    cov = json.loads(COVERAGE.read_text()) if COVERAGE.exists() else {"states": {}}
    entry = cov["states"].setdefault(state, {"name": state, "collected": {}})
    entry.setdefault("collected", {})[key] = info
    if entry.get("status") in (None, "pending", "researching"):
        entry["status"] = "collected"
    cov["updated_at"] = datetime.now(timezone.utc).isoformat()
    COVERAGE.write_text(json.dumps(cov, indent=2))


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--state", nargs="+", help="limit to these state codes")
    ap.add_argument("-v", "--verbose", action="store_true")
    args = ap.parse_args()

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s", datefmt="%H:%M:%S")

    if not SOURCES.exists():
        print(f"{SOURCES} not found — nothing configured yet", file=sys.stderr)
        return 1
    config = yaml.safe_load(SOURCES.read_text()) or {}

    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT

    corps_fresh: list[dict] = []
    regs_fresh: list[dict] = []
    corps_states: set[str] = set()
    regs_states: set[str] = set()

    for state, sources in sorted(config.items()):
        if args.state and state not in args.state:
            continue
        for src in sources or []:
            kind = src["kind"]
            label = f"{state} {kind} ({src.get('source', src.get('dataset', src.get('url', '?')))})"
            started = time.time()
            try:
                if kind == "socrata_corp":
                    rows = records_socrata.fetch(session, src, state)
                    if src.get("dedupe"):
                        seen: set = set()
                        rows = [r for r in rows
                                if r["record_id"] not in seen
                                and not seen.add(r["record_id"])]
                    corps_fresh.extend(rows)
                    corps_states.add(state)
                    update_coverage(state, "corporate", {
                        "source": src.get("source", ""), "records": len(rows),
                        "url": src.get("page", ""),
                        "coverage": src.get("coverage", "all")})
                elif kind == "socrata_registry":
                    rows = records_socrata.fetch(session, src, state,
                                                 filter_names=src.get("filter_names", False))
                    regs_fresh.extend(rows)
                    regs_states.add(state)
                    update_coverage(state, "hoa_registry", {
                        "source": src.get("source", ""), "records": len(rows),
                        "url": src.get("page", "")})
                elif kind == "tecuity_corp":
                    rows = records_tecuity.fetch(session, src, state)
                    corps_fresh.extend(rows)
                    corps_states.add(state)
                    update_coverage(state, "corporate", {
                        "source": src.get("source", ""), "records": len(rows),
                        "url": src.get("page", ""),
                        "coverage": "partial"})
                elif kind == "arcgis_corp":
                    rows = fetch_arcgis(session, src, state)
                    corps_fresh.extend(rows)
                    corps_states.add(state)
                    update_coverage(state, "corporate", {
                        "source": src.get("source", ""), "records": len(rows),
                        "url": src.get("page", ""),
                        "coverage": src.get("coverage", "all")})
                elif kind == "csv_corp":
                    rows = fetch_csv(session, src, state, registry=False)
                    corps_fresh.extend(rows)
                    corps_states.add(state)
                    update_coverage(state, "corporate", {
                        "source": src.get("source", ""), "records": len(rows),
                        "url": src.get("page", ""),
                        "coverage": src.get("coverage", "all")})
                elif kind == "csv_registry":
                    rows = fetch_csv(session, src, state, registry=True)
                    regs_fresh.extend(rows)
                    regs_states.add(state)
                    update_coverage(state, "hoa_registry", {
                        "source": src.get("source", ""), "records": len(rows),
                        "url": src.get("page", "")})
                else:
                    log.warning("%s: unknown kind %r", state, kind)
                    continue
                log.info("%s -> %d rows in %.0fs", label, len(rows), time.time() - started)
            except Exception as exc:
                log.error("%s FAILED: %s: %s", label, type(exc).__name__, exc)

    # Coverage semantics per state, for the site build: "active_only" states
    # support absence-of-corporation flags, "partial" states never do.
    cov_map = {}
    for state, sources in config.items():
        for src in sources or []:
            if src["kind"].endswith("_corp"):
                cov_map[state] = ("partial" if src["kind"] == "tecuity_corp"
                                  else src.get("coverage", "all"))
    (OUT_DIR / "corp_coverage.json").write_text(json.dumps(cov_map, indent=2))

    if corps_states:
        total = write_merged(CORPS, corps_fresh, corps_states)
        log.info("state_corps.jsonl: %d rows total (%d fresh from %s)",
                 total, len(corps_fresh), ", ".join(sorted(corps_states)))
    if regs_states:
        total = write_merged(REGS, regs_fresh, regs_states)
        log.info("state_registries.jsonl: %d rows total (%d fresh from %s)",
                 total, len(regs_fresh), ", ".join(sorted(regs_states)))

    by_state = Counter(r["state"] for r in corps_fresh)
    for st, n in sorted(by_state.items()):
        print(f"  {st}: {n:,} association-named corporations")
    by_state = Counter(r["state"] for r in regs_fresh)
    for st, n in sorted(by_state.items()):
        print(f"  {st}: {n:,} registry records")
    return 0


if __name__ == "__main__":
    sys.exit(main())

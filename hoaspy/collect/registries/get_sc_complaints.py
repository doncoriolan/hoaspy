#!/usr/bin/env python
"""South Carolina HOA complaint data from the Department of Consumer Affairs.

S.C. Code 37-6-117 makes SCDCA collect HOA complaints and report them to the
legislature annually — the reports name the association, its county, its
management company, the complaint category, and the outcome. The XLSX links
are scraped off the HOA-reports page so new years appear automatically.

    ./venv/bin/python -m hoaspy.collect.registries.get_sc_complaints

Output: records/sc_complaints.jsonl — one row per complaint.
"""

from __future__ import annotations

import io
import json
import logging
import re
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import requests

from hoaspy import ROOT
OUT = ROOT / "records" / "sc_complaints.jsonl"
PAGE = "https://consumer.sc.gov/HOA-reports"
BASE = "https://consumer.sc.gov"

USER_AGENT = "macos:hoa-public-records:0.1 (academic research)"

log = logging.getLogger("sc")

# The header wording drifts a little between years; match loosely.
COLS = {
    "number": r"complaint number",
    "date": r"complaint date",
    "name": r"hoa name",
    "city": r"hoa city",
    "county": r"hoa county",
    "manager": r"management company name",
    "category": r"^complaint description$",
    "status": r"^status$",
}


def clean(v) -> str:
    s = str(v) if v is not None else ""
    return " ".join(s.replace("_x000D_", " ").split())


def parse_book(content: bytes, url: str) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    out = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        try:
            header = [clean(h).lower() for h in next(rows)]
        except StopIteration:
            continue
        idx = {}
        for key, pat in COLS.items():
            for i, h in enumerate(header):
                if re.search(pat, h):
                    idx[key] = i
                    break
        if "name" not in idx:
            continue

        def cell(row, key):
            i = idx.get(key)
            return clean(row[i]) if i is not None and i < len(row) else ""

        for row in rows:
            name = cell(row, "name")
            if not name or name.lower() in ("none", "n/a", "unknown"):
                continue
            date = cell(row, "date")[:10]
            number = cell(row, "number")
            year = None
            if m := re.search(r"\b(20[12]\d)\b", date):
                year = int(m.group(1))
            elif m := re.match(r"C(\d{2})-", number):
                year = 2000 + int(m.group(1))
            out.append({
                "state": "SC",
                "source": "sc-dca-hoa-complaints",
                "source_url": url,
                "complaint_number": number,
                "date": date,
                "year": year,
                "name": name,
                "city": cell(row, "city").title(),
                "county": cell(row, "county").title(),
                "manager_name": cell(row, "manager"),
                "category": cell(row, "category"),
                "status": cell(row, "status"),
            })
    return out


def main() -> int:
    logging.basicConfig(level=logging.INFO,
                        format="%(asctime)s %(levelname)-7s %(message)s",
                        datefmt="%H:%M:%S")
    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    started = time.time()

    page = session.get(PAGE, timeout=60)
    page.raise_for_status()
    links = sorted(set(re.findall(r'href="([^"]*\.xlsx)"', page.text, re.I)))
    log.info("found %d report workbooks", len(links))

    records: list[dict] = []
    for href in links:
        url = href if href.startswith("http") else BASE + href
        try:
            resp = session.get(url, timeout=120)
            resp.raise_for_status()
            rows = parse_book(resp.content, url)
            records.extend(rows)
            log.info("%s -> %d complaints", url.rsplit("/", 1)[-1], len(rows))
        except Exception as exc:
            log.error("%s failed: %s", url, exc)
        time.sleep(1.0)

    # A complaint can appear in two adjacent reports; dedupe on number+name.
    seen, deduped = set(), []
    for r in records:
        key = (r["complaint_number"], r["name"].lower())
        if key in seen:
            continue
        seen.add(key)
        deduped.append(r)

    now = datetime.now(timezone.utc).isoformat()
    for r in deduped:
        r["retrieved_at"] = now
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w") as fh:
        for r in deduped:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")

    by_year = Counter(r["year"] for r in deduped if r["year"])
    print(f"\n{len(deduped):,} complaints ({len(records) - len(deduped)} duplicates dropped)")
    for y in sorted(by_year):
        print(f"  {y}: {by_year[y]}")
    print(f"Wrote {OUT} in {time.time() - started:.0f}s")
    return 0


if __name__ == "__main__":
    sys.exit(main())
